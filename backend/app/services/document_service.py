import csv
import io
import os
import re
import uuid
from pathlib import Path

from sqlalchemy import text
from sqlalchemy.orm import Session

from app.config import get_storage_root
from app.models import Document

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv"}
MAX_FILE_SIZE = 50 * 1024 * 1024  # 50MB


def _client_doc_root(client_id: str) -> Path:
    return get_storage_root() / "clients" / client_id / "documents"


def _client_extracted_root(client_id: str) -> Path:
    return get_storage_root() / "clients" / client_id / "extracted"


def _sanitize_filename(filename: str) -> str:
    name = re.sub(r'[<>:"/\\|?*]', "_", filename)
    name = re.sub(r"\.{2,}", ".", name)
    name = name.strip(". ")
    if not name:
        name = "unnamed"
    return name[:255]


def _extract_pdf(file_path: str) -> str:
    from PyPDF2 import PdfReader

    reader = PdfReader(file_path)
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
    return "\n".join(text_parts)


def _extract_docx(file_path: str) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(file_path)
    return "\n".join(para.text for para in doc.paragraphs if para.text)


def _extract_csv(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f)
        rows = list(reader)
    return "\n".join(", ".join(row) for row in rows)


def _extract_text(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def extract_text_from_file(file_path: str, mime_type: str) -> str:
    if mime_type == "application/pdf":
        return _extract_pdf(file_path)
    elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _extract_docx(file_path)
    elif mime_type == "text/csv":
        return _extract_csv(file_path)
    else:
        return _extract_text(file_path)


def get_mime_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    mime_map = {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".csv": "text/csv",
    }
    return mime_map.get(ext, "application/octet-stream")


def validate_file(filename: str, file_size: int) -> tuple[bool, str | None]:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        return False, f"Unsupported file type: {ext}. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
    if file_size > MAX_FILE_SIZE:
        return False, f"File too large: {file_size} bytes. Maximum: {MAX_FILE_SIZE} bytes"
    return True, None


def upload_document(
    db: Session,
    client_id: str,
    filename: str,
    file_content: bytes,
    title: str | None = None,
) -> Document:
    safe_name = _sanitize_filename(filename)
    doc_id = str(uuid.uuid4())
    stored_name = f"{doc_id}_{safe_name}"

    doc_dir = _client_doc_root(client_id)
    doc_dir.mkdir(parents=True, exist_ok=True)
    file_path = doc_dir / stored_name
    file_path.write_bytes(file_content)

    mime_type = get_mime_type(safe_name)
    file_size = len(file_content)

    doc = Document(
        id=doc_id,
        client_id=client_id,
        title=title or safe_name,
        file_path=str(file_path),
        mime_type=mime_type,
        file_size=file_size,
        extraction_status="pending",
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    try:
        extracted = extract_text_from_file(str(file_path), mime_type)
        extracted_dir = _client_extracted_root(client_id)
        extracted_dir.mkdir(parents=True, exist_ok=True)
        extracted_path = extracted_dir / f"{doc_id}.txt"
        extracted_path.write_text(extracted, encoding="utf-8")

        with db.bind.connect() as conn:
            conn.execute(
                text("INSERT INTO document_content (document_id, client_id, content) VALUES (:doc_id, :client_id, :content)"),
                {"doc_id": doc_id, "client_id": client_id, "content": extracted},
            )
            conn.commit()

        doc.extraction_status = "completed"
    except Exception as e:
        doc.extraction_status = "failed"
        doc.extraction_error = str(e)

    db.commit()
    db.refresh(doc)
    return doc


def list_documents(db: Session, client_id: str) -> list[Document]:
    return db.query(Document).filter(Document.client_id == client_id).order_by(Document.created_at.desc()).all()
